"""시험판 registry 만들기 — prefill 판(paths.REGISTRY_DIR) 파일 전부 + mapping/registry_v<판>.csv 를 KRICT-PSC 에 얹어 <출력>/v<판>/ 에.

정본 아님(이 레포는 prefill registry 를 고치지 않는다 — CLAUDE.md). 쓰기: 빌드 때 SDL_REGISTRY_DIR=<출력>/v<판> (prefill paths 고리).
사용: python3 scripts/registry_trial.py <$SOLAR_DATA_DIR/registry> 0.5.4b
"""
import csv, hashlib, os, pathlib, shutil, subprocess, sys

HERE = pathlib.Path(__file__).resolve().parent
PREFILL = (HERE.parent / os.environ.get("SDL_PREFILL_DIR", "../sdl-prefill")).resolve()
assert not os.environ.get("SDL_REGISTRY_DIR"), "원본 판에서 만든다 — SDL_REGISTRY_DIR 을 풀고 실행"
sys.path[:0] = [str(PREFILL)]
import openpyxl  # noqa: E402
from sdl_platform import paths  # noqa: E402
from sdl_platform.engine import lab_registry  # noqa: E402

LAB = "KRICT-PSC"


def main(out_root, ver):
    patch = HERE.parent / "mapping" / f"registry_v{ver}.csv"
    rows = list(csv.DictReader(line for line in patch.open(encoding="utf-8") if not line.startswith("#")))
    out = pathlib.Path(out_root) / f"v{ver}"
    if out.exists():
        shutil.rmtree(out)                             # 만든 것(다시 만들 수 있음)만 지운다
    out.mkdir(parents=True)
    src = sorted(pathlib.Path(paths.REGISTRY_DIR).glob("sdl_registry_v*_*.xlsx"))
    for f in src:
        shutil.copy2(f, out / f.name.replace(f"v{paths.REGISTRY_VERSION}_", f"v{ver}_"))
    psc = next(out.glob(f"*_{LAB}.xlsx"))
    wb = openpyxl.load_workbook(psc)
    tag = f"v{ver}: 시험판 — sdl-solar-chemdx-library mapping/registry_v{ver}.csv(ChemDX 건너뛴 레코드)"
    for r in rows:
        ws = wb[r["sheet"]]
        hr = next(i for i in range(1, 12) if ws.cell(i, 1).value == "컬럼명")
        col = {ws.cell(hr, c).value: c for c in range(2, ws.max_column + 1) if ws.cell(hr, c).value}
        at = max(i for i in range(hr, ws.max_row + 1) if any(ws.cell(i, c).value not in (None, "") for c in col.values())) + 1
        ws.cell(at, 1).value = "추가"
        for k, v in r.items():
            if k in col and v:
                ws.cell(at, col[k]).value = v
        ws.cell(at, col["수정이력"]).value = tag
    wb.save(psc)
    data = lab_registry._parse_v2_file(str(psc))     # 검증 게이트 = prefill 과 같은 함수
    assert lab_registry._validate(str(psc), data) == LAB
    codes = {x["StepTypeCode"] for x in data["StepType"]}
    assert {r["StepTypeCode"] for r in rows if r["sheet"] == "StepType"} <= codes
    commit = subprocess.run(["git", "describe", "--always", "--dirty"], cwd=PREFILL, capture_output=True, text=True).stdout.strip()
    (out / "SOURCE.txt").write_text(
        f"시험판 registry v{ver} — 정본 아님\nprefill {commit} · 원본 판 v{paths.REGISTRY_VERSION}\n"
        f"패치 mapping/{patch.name} sha256 {hashlib.sha256(patch.read_bytes()).hexdigest()[:16]} · 행 {len(rows)}\n", encoding="utf-8")
    print(f"{out} — 파일 {len(src)} · {LAB} 행 +{len(rows)} · StepType {len(codes)}")


if __name__ == "__main__":
    main(*sys.argv[1:3])
